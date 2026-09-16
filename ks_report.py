#!/usr/bin/env python3
"""Таблица разбора отчёта kubescape (CIS) по выжимке ks_sanitize.py.

Запуск:
  python ks_report.py cis-v1.12.0-proxy.sanitized.json
  python ks_report.py cis-v1.12.0-proxy.sanitized.json --names cis-v1.12.0-proxy.users.csv

Рядом появится cis-v1.12.0-proxy.analysis.xlsx. С --names псевдонимы user-NN
заменяются настоящими именами (cis-v1.12.0-proxy.analysis-names.xlsx):
этот вариант — для передачи командам внутри компании.
"""

import argparse
import csv
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# namespace платформы: их чинит команда кластера, остальное — команды сервисов
PLATFORM_NS = {
    "kube-system", "kube-public", "kube-node-lease", "default", "argocd", "cert-manager",
    "ceph-csi-cephfs", "ceph-csi-rbd", "cilium-secrets", "consul", "devops", "external-secrets",
    "gitlab-runner", "gradle-cache", "ingress-nginx", "jaeger", "kubernetes-dashboard", "logging",
    "monitoring", "nfs", "prometheus", "rabbitmq-operator", "vault",
}
SUBJECT_KINDS = {"User", "Group", "ServiceAccount"}
RBAC_KINDS = SUBJECT_KINDS | {"Role", "ClusterRole", "RoleBinding", "ClusterRoleBinding"}
SEVERITY_RU = {"Critical": "критическая", "High": "высокая", "Medium": "средняя", "Low": "низкая"}
SEVERITY_ORDER = {"критическая": 0, "высокая": 1, "средняя": 2, "низкая": 3}
STATUS_RU = {"failed": "не пройдена", "passed": "пройдена", "skipped": "не проверялась"}
STATUS_ORDER = {"не пройдена": 0, "не проверялась": 1, "пройдена": 2}
PRIORITY_ORDER = {"P1": 0, "P2": 1, "P3": 2, "инфо": 3, "—": 4}
PRIORITY_FILL = {"P1": "F4B6B6", "P2": "F9D9AE", "P3": "FFF1B8", "инфо": "E2E2E2"}
PSEUDO = re.compile(r"\b(?:user|email)-\d{2,}\b")
ENV_MASKED = "переменные-секреты, значение скрыто сканером"  # вид места из ks_sanitize.py
FIX_NOISE = re.compile(r"seLinuxOptions|fsGroupChangePolicy")
CHECKS = "'Проверки'"

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
TITLE = Font(name="Arial", size=14, bold=True)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="DDEBF7")

PSA_WHAT = "В namespace не включён Pod Security Admission, поэтому ничто не мешает запустить такой под."
PSA_FIX = ("Метка pod-security.kubernetes.io/enforce={level} на namespace. Сначала warn и audit, "
           "чтобы найти нарушителей: kubectl label --dry-run=server --overwrite ns --all "
           "pod-security.kubernetes.io/enforce={level}")


def psa(level):
    return ("P1", "платформа", PSA_WHAT, PSA_FIX.format(level=level))


# приоритет, кто исправляет, что не так, что сделать — для непройденных проверок
CONTROL_NOTES = {
    "C-0211": ("P2", "команды сервисов",
               "У подов и контейнеров не задан securityContext: запуск не от root, запрет повышения "
               "привилегий, файловая система только для чтения, сброс capabilities.",
               "Задать в общих шаблонах Helm-чартов: runAsNonRoot: true, allowPrivilegeEscalation: false, "
               "readOnlyRootFilesystem: true (где приложение позволяет), capabilities.drop: [ALL]. "
               "seLinuxOptions — только если на узлах включён SELinux."),
    "C-0193": psa("baseline"),
    "C-0197": psa("restricted"),
    "C-0198": psa("restricted"),
    "C-0199": psa("restricted"),
    "C-0200": psa("baseline"),
    "C-0201": psa("restricted"),
    "C-0203": psa("baseline"),
    "C-0204": psa("baseline"),
    "C-0202": ("инфо", "платформа",
               "Windows-контейнеры с доступом к узлу. Если узлов Windows нет, риска нет.",
               "Закроется той же меткой Pod Security Admission (baseline)."),
    "C-0041": ("P3", "платформа",
               "Поды с hostNetwork — общая сеть с узлом. Обычно это системные компоненты "
               "(сетевой плагин, ingress, мониторинг), которым это нужно по устройству.",
               "Проверить по листу «Находки», что в списке только системные компоненты; "
               "прикладным подам hostNetwork не выдавать. Запрет для прикладных namespace даёт "
               "та же метка Pod Security Admission (baseline)."),
    "C-0275": ("P3", "платформа",
               "Поды с hostPID видят процессы узла.",
               "То же, что и с hostNetwork: убедиться, что это только системные компоненты; "
               "для остальных namespace запрет даёт метка Pod Security Admission (baseline)."),
    "C-0187": ("P3", "платформа",
               "Роли со звёздочкой в правах: доступ ко всем ресурсам или ко всем действиям.",
               "Заменить * на явный список ресурсов и действий; для ролей операторов свериться "
               "с их документацией. Список — лист «RBAC»."),
    "C-0185": ("P1", "платформа",
               "Роль cluster-admin — полный доступ ко всему кластеру, включая все Secret — выдана "
               "не только аварийной учётной записи.",
               "Людям — вход через OIDC с ограниченными ролями; cluster-admin оставить аварийной "
               "учётке; сервисным аккаунтам — минимальные роли. Список — лист «RBAC»."),
    "C-0117": ("P2", "платформа",
               "API-сервер не проверяет сертификат kubelet, когда обращается к узлу (exec, logs, "
               "port-forward), поэтому возможен перехват.",
               "Включить серверные сертификаты kubelet от CA кластера (в Kubespray — "
               "kubelet_rotate_server_certificates) и задать --kubelet-certificate-authority."),
    "C-0190": ("P2", "команды сервисов",
               "В под монтируется токен сервисного аккаунта, хотя большинству приложений API "
               "Kubernetes не нужен. При взломе контейнера токен достаётся атакующему.",
               "automountServiceAccountToken: false в шаблонах; токен оставить только тем, кто "
               "ходит в API (операторы, CI)."),
    "C-0210": ("P2", "команды сервисов",
               "Не задан профиль seccomp — фильтр системных вызовов контейнера. Название проверки "
               "устарело: сейчас это RuntimeDefault, а не docker/default.",
               "seccompProfile.type: RuntimeDefault в securityContext пода. Можно включить сразу для "
               "всего кластера (seccompDefault в kubelet), но kubescape смотрит на манифесты и этого "
               "не увидит. Внедрять по namespace: изредка приложения с ним падают."),
    "C-0189": ("P3", "команды сервисов",
               "Поды работают под сервисным аккаунтом default, поэтому у всех подов namespace "
               "одни и те же права.",
               "Отдельный сервисный аккаунт на приложение; у default — automountServiceAccountToken: false."),
    "C-0209": ("инфо", "—",
               "Ручной пункт CIS: kubescape не может сам оценить, как ресурсы поделены по namespace, "
               "и выносит все namespace на просмотр.",
               "Работ не требует."),
    "C-0206": ("P2", "платформа, команды сервисов",
               "В namespace нет сетевых политик, поэтому любой под может обратиться к любому.",
               "Начать с запрета входящего трафика по умолчанию и явных разрешений, в первую "
               "очередь в namespace с ПДн. Если сетевой плагин Cilium, сначала проверить политики "
               "CiliumNetworkPolicy (kubescape их не учитывает): kubectl get cnp,ccnp -A."),
    "C-0186": ("P2", "платформа",
               "Субъекты, которым можно читать Secret.",
               "Проверить по листу «RBAC»: операторам (cert-manager, external-secrets) это нужно, "
               "людям и CI — как правило, нет."),
    "C-0188": ("P3", "платформа",
               "Право создавать поды: можно запустить под с любым сервисным аккаунтом namespace "
               "и получить его права.",
               "Оставить только CI/CD и операторам."),
    "C-0279": ("P2", "платформа",
               "Доступ к nodes/proxy — прямой вызов API kubelet: можно выполнять команды в подах "
               "узла в обход журнала аудита API-сервера.",
               "Убрать у всех, кому это не нужно для работы (например, оставить мониторингу)."),
    "C-0281": ("P3", "платформа",
               "Право менять admission-вебхуки: можно отключить проверки или перехватывать "
               "создаваемые объекты.",
               "Оставить только операторам, которые ставят свои вебхуки."),
    "C-0282": ("P3", "платформа",
               "Право выпускать токены сервисных аккаунтов: можно получить токен любого аккаунта namespace.",
               "Оставить только тем, кому это нужно."),
    "C-0278": ("P3", "платформа",
               "Право создавать PersistentVolume: можно создать том на hostPath и добраться до файлов узла.",
               "Оставить только CSI-драйверам и администраторам."),
    "C-0280": ("P3", "платформа",
               "Право одобрять запросы на сертификаты: можно выпустить клиентский сертификат на "
               "любое имя, в том числе администратора.",
               "Оставить только controller-manager и одобрителю сертификатов kubelet."),
    "C-0191": ("P3", "платформа",
               "Права bind, escalate, impersonate: позволяют выдать себе чужие права.",
               "Оставить только администраторам кластера."),
    "C-0207": ("P3", "команды сервисов",
               "Приложение получает Secret через переменные окружения, а не файлом: значения видны "
               "в /proc, дампах и иногда в логах.",
               "Монтировать Secret файлом."),
    "C-0113": ("P1", "платформа",
               "API-сервер принимает запросы без аутентификации: анонимный запрос приходит от "
               "system:anonymous.",
               "--anonymous-auth=false. Заранее проверить, что проверки живости узлов и "
               "внешние клиенты ходят с сертификатом или токеном."),
    "C-0129": ("P3", "платформа",
               "У API-сервера включён профайлер (/debug/pprof).",
               "--profiling=false; включать только на время разбора проблем "
               "с производительностью."),
    "C-0130": ("P1", "платформа",
               "Журнал аудита API-сервера не ведётся: не задан --audit-log-path.",
               "Задать --audit-log-path и --audit-policy-file, журнал складывать в "
               "централизованное хранилище."),
    "C-0131": ("P3", "платформа",
               "Срок хранения журнала аудита меньше 30 дней.",
               "--audit-log-maxage не меньше 30; для 152-ФЗ срок согласовать с требованиями "
               "к хранению событий безопасности."),
    "C-0141": ("P1", "платформа",
               "Содержимое etcd не шифруется: не задан --encryption-provider-config, значит все "
               "Secret лежат в базе кластера в открытом виде.",
               "Файл EncryptionConfiguration с провайдером (лучше KMS, иначе aescbc/secretbox) и "
               "--encryption-provider-config; после включения перезаписать существующие секреты: "
               "kubectl get secrets -A -o json | kubectl replace -f -"),
    "C-0160": ("P1", "платформа",
               "Нет политики аудита: не задан --audit-policy-file, поэтому события API-сервера "
               "не отбираются и не пишутся.",
               "Создать политику аудита (за основу — пример из документации Kubernetes), "
               "включить её вместе с --audit-log-path."),
    "C-0291": ("P3", "платформа",
               "Метрики kube-proxy слушают не только localhost.",
               "metricsBindAddress: 127.0.0.1:10249 в конфигурации kube-proxy; если метрики "
               "собирает Prometheus, оставить доступ только ему сетевой политикой."),
    "C-0121": ("P3", "платформа",
               "Не включён admission-плагин EventRateLimit, ограничивающий поток событий к API-серверу.",
               "Добавить в --enable-admission-plugins с файлом настроек."),
    "C-0123": ("P3", "платформа",
               "Не включён AlwaysPullImages: под может запуститься из приватного образа, уже "
               "скачанного на узел для другого namespace.",
               "Включить, если на узлах работают команды с разными правами на реестр. "
               "Вырастет нагрузка на реестр."),
    "C-0132": ("P3", "платформа",
               "Журнал аудита API-сервера включён, но хранит меньше 10 архивов.",
               "--audit-log-maxbackup не меньше 10 или отправка аудита в централизованное "
               "хранилище с нужным сроком хранения."),
    "C-0133": ("P3", "платформа",
               "Размер файла журнала аудита меньше 100 МБ.",
               "--audit-log-maxsize не меньше 100."),
    "C-0134": ("P3", "платформа",
               "Значение --request-timeout отличается от рекомендованного.",
               "Проверить вручную: CIS допускает любое обоснованное значение."),
    "C-0277": ("P2", "платформа",
               "Не ограничен список шифров TLS API-сервера.",
               "--tls-cipher-suites: только ECDHE с AES-GCM или ChaCha20. Значение из отчёта "
               "kubescape не копировать: в нём есть RC4 и CBC."),
    "C-0283": ("P2", "платформа",
               "Не включён плагин DenyServiceExternalIPs: можно создать Service с чужим externalIP "
               "и перехватить трафик (CVE-2020-8554).",
               "Добавить DenyServiceExternalIPs в --enable-admission-plugins."),
    "C-0290": ("P3", "платформа",
               "API-сервер продлевает токены сервисных аккаунтов до года ради старых клиентов.",
               "--service-account-extend-token-expiration=false; перед этим убедиться, что "
               "приложения, работающие с API, перечитывают токен."),
    "C-0212": ("инфо", "—",
               "Находки в namespace default. Если это только служебный EndpointSlice kubernetes — "
               "ложное срабатывание.",
               "Проверить список на листе «Находки»."),
}
SKIPPED_NOTE = ("—", "платформа",
                "kubescape в режиме CLI этого не проверяет: нужен доступ к файлам и настройкам на узлах.",
                "Разово запустить kube-bench на узлах или поставить оператор kubescape.")
FIX_OVERRIDE = {"C-0277": "задать --tls-cipher-suites без RC4 и CBC (см. лист «Проверки»)"}


def cis_parts(name):
    m = re.match(r"CIS-([\d.]+)\s+(.*)", name)
    return (m.group(1), m.group(2)) if m else ("", name)


def zone_of(ns):
    if not ns:
        return "кластер"
    return "платформа" if ns in PLATFORM_NS else "сервисы"


def check_value(col, cid):
    """Формула: значение столбца col листа «Проверки» для проверки cid."""
    return (f"IFERROR(INDEX({CHECKS}!${col}$2:${col}$500,"
            f"MATCH(\"{cid}\",{CHECKS}!$A$2:$A$500,0)),0)")


def failed(cid):
    return check_value("G", cid)


def total(cid):
    return f"({check_value('G', cid)}+{check_value('H', cid)})"


def fix_text(f):
    if f["control"] in FIX_OVERRIDE:
        return FIX_OVERRIDE[f["control"]]
    if f["kind"] in SUBJECT_KINDS and f.get("role"):
        return "сузить роль или убрать привязку"
    p = f.get("paths", {})
    parts = [f"{x['path']} = {x['value']}" for x in p.get("fix", [])
             if x["value"] != "YOUR_VALUE" and not FIX_NOISE.search(x["path"])]
    parts += [f"удалить: {x}" for x in p.get("delete", [])]
    parts += [f"проверить: {x}" for x in p.get("review", [])]
    if not parts:
        parts = [f"поле: {x}" for x in p.get("failed", [])]
    text = "; ".join(dict.fromkeys(parts))
    return text if len(text) <= 600 else text[:600] + "…"


def whose(f):
    name, kind = f["name"], f["kind"]
    if PSEUDO.search(name):
        return "человек"
    if name.startswith(("system:", "kubeadm:")):
        return "система"
    if f["namespace"] in PLATFORM_NS or kind in ("ClusterRole", "Role"):
        return "платформа или оператор"
    return "приложение или CI"


def load_names(path):
    with open(path, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh, delimiter=";"))
    return {r[0]: r[1] for r in rows[1:] if len(r) >= 2}


def load_secret_places(path):
    if not path.exists():
        return []
    with path.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh, delimiter=";"))
    places = []
    for kind_of_place, where in (r for r in rows[1:] if len(r) >= 2):
        head, _, detail = where.partition(" · ")
        kind, _, nsname = head.partition(" ")
        ns, _, name = nsname.partition("/")
        ns = "" if ns == "-" else ns
        places.append((kind_of_place, ns, kind, name, detail))
    return places


SECRET_ACTION = {
    ENV_MASKED: "Значение в отчёте скрыто сканером, но в манифесте оно задано строкой, а не ссылкой на "
                "Secret. Проверить в чарте или репозитории: если там настоящий секрет — перенести в Secret "
                "(лучше через External Secrets из Vault) и сменить.",
    "переменные окружения": "Перенести значение в Secret (лучше через External Secrets из Vault), "
                            "подключать через secretKeyRef или файлом; значение сменить.",
    "аргументы запуска": "Не передавать секрет в аргументах: он виден в описании пода и в списке "
                         "процессов. Читать из файла или из переменной из Secret; значение сменить.",
    "пароли внутри значений": "Внутри значения есть пара «password: …» или логин с паролем в адресе. "
                              "Проверить и вынести в Secret.",
    "ключи ConfigMap": "ConfigMap не для секретов: перенести в Secret; значение сменить.",
    "поля объектов": "Поле с именем как у секрета: проверить.",
}


class Sheet:
    """Лист с таблицей: шапка, данные, ширины, фильтр, закреплённая строка."""

    def __init__(self, wb, title, headers, widths, wrap=(), priority_col=None, restore=None):
        self.ws = wb.create_sheet(title)
        self.headers, self.widths = headers, widths
        self.wrap, self.priority_col = set(wrap), priority_col
        self.restore = restore or (lambda s: s)
        self.ws.append(headers)

    def add(self, row):
        self.ws.append([self.restore(v) if isinstance(v, str) and not v.startswith("=") else v
                        for v in row])

    def finish(self):
        ws = self.ws
        for i, w in enumerate(self.widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for cell in ws[1]:
            cell.font, cell.fill = HEADER_FONT, HEADER_FILL
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = FONT
                cell.alignment = Alignment(vertical="top", wrap_text=cell.column in self.wrap)
            if self.priority_col:
                cell = row[self.priority_col - 1]
                if cell.value in PRIORITY_FILL:
                    cell.fill = PatternFill("solid", fgColor=PRIORITY_FILL[cell.value])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(self.headers))}{max(ws.max_row, 2)}"
        fit_to_width(ws, landscape=True)
        ws.print_title_rows = "1:1"


def fit_to_width(ws, landscape=False):
    """При печати лист помещается по ширине страницы."""
    if landscape:
        ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def main():
    ap = argparse.ArgumentParser(description="Таблица разбора отчёта kubescape по выжимке.")
    ap.add_argument("sanitized", help="*.sanitized.json от ks_sanitize.py")
    ap.add_argument("--names", help="*.users.csv от ks_sanitize.py: вернуть настоящие имена")
    args = ap.parse_args()
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(encoding="utf-8", errors="replace")

    src = Path(args.sanitized).resolve()
    data = json.loads(src.read_text(encoding="utf-8"))
    base = src.name.removesuffix(".json").removesuffix(".sanitized")
    dst = src.with_name(base + (".analysis-names.xlsx" if args.names else ".analysis.xlsx"))
    names = load_names(args.names) if args.names else {}

    def restore(s):
        return PSEUDO.sub(lambda m: names.get(m.group(0), m.group(0)), s) if names else s

    findings = data["findings"]
    failed_findings = [f for f in findings if f["status"] == "failed"]
    namespaces = {f["name"] for f in findings if f["kind"] == "Namespace"}
    ns_seen = namespaces | {f["namespace"] for f in findings if f["namespace"]}

    # что можно понять о кластере по самим находкам
    def names_of(cid):
        return sorted({f["name"] for f in failed_findings if f["control"] == cid})

    np_bad = set(names_of("C-0206"))
    np_excepted = {f["name"] for f in findings if f["control"] == "C-0206" and f.get("subStatus")}
    np_ok = sorted(namespaces - np_bad - np_excepted)
    host_net = [f for f in failed_findings if f["control"] in ("C-0041", "C-0275")]
    host_all_platform = bool(host_net) and all(f["namespace"] in PLATFORM_NS for f in host_net)
    cilium = any("cilium" in f["name"] for f in findings) or "cilium-secrets" in ns_seen
    kubespray = any(f["name"].startswith("nginx-proxy-") and f["namespace"] == "kube-system"
                    for f in findings)
    has_vault, has_eso = "vault" in ns_seen, "external-secrets" in ns_seen
    default_ns = [f for f in failed_findings if f["control"] == "C-0212"]

    notes = dict(CONTROL_NOTES)
    prio, owner, what, fix = notes["C-0206"]
    what += f" Политики есть только в: {', '.join(np_ok)}." if np_ok else " Политик нет ни в одном namespace."
    if cilium:
        what += " Сетевой плагин — Cilium, политики он поддерживает."
    notes["C-0206"] = (prio, owner, what, fix)
    for cid, word in (("C-0041", "сеть"), ("C-0275", "процессы")):
        hosts = sorted({f"{f['namespace']}/{f['name']}" for f in host_net if f["control"] == cid})
        if not hosts:
            continue
        if host_all_platform:
            notes[cid] = ("инфо", "платформа",
                          f"Под использует {word} узла. Все такие поды системные"
                          + (" (nginx-proxy — так Kubespray балансирует доступ к API)" if kubespray else "")
                          + ": им это нужно.",
                          "Оформить исключения в kubescape, чтобы не мешали в следующих сканах.")
        else:
            notes[cid] = ("P2", "платформа, команды сервисов",
                          f"Под использует {word} узла: при взломе — прямой путь на узел.",
                          "Убрать hostNetwork/hostPID у прикладных подов; системным оформить исключения.")
    if default_ns and all(f["kind"] == "EndpointSlice" and f["name"] == "kubernetes" for f in default_ns):
        notes["C-0212"] = ("инфо", "—",
                           "Единственная находка — служебный EndpointSlice kubernetes. Ложное срабатывание.",
                           "Работ не требует.")
    skipped_notes = {
        "C-0205": ("—", "платформа",
                   "Не проверялась." + (" Сетевой плагин — Cilium (видно по подам), политики он поддерживает."
                                        if cilium else " Выяснить, какой сетевой плагин и поддерживает ли он политики."),
                   "Работ не требует." if cilium else "Проверить вручную."),
        "C-0208": ("—", "платформа",
                   "Не проверялась." + (" В кластере есть External Secrets и Vault." if has_vault and has_eso else ""),
                   "Хранить секреты во внешнем хранилище: см. лист «Секреты»."),
    }

    wb = Workbook()
    summary = wb.active
    summary.title = "Сводка"

    # --- План: корневые причины вместо сотни отдельных проверок
    secret_places = load_secret_places(src.with_name(base + ".secret-places.csv"))
    examples = []
    for kind_of_place, _, _, _, detail in secret_places:
        var = detail.rsplit(" · ", 1)[-1]
        if (kind_of_place in ("переменные окружения", ENV_MASKED)
                and re.search(r"ADMIN|DB_PASS|JWT|R2DBC|MASTER", var)):
            examples.append(var)
    examples = list(dict.fromkeys(examples))[:5]
    store = ("Vault через External Secrets (оба уже стоят в кластере)" if has_vault and has_eso
             else "внешнее хранилище секретов")
    masked_env = any(p[0] == ENV_MASKED for p in secret_places)
    open_places = [p for p in secret_places if p[0] != ENV_MASKED]
    risk = ("Всё, что записано в описании пода, видит любой, кто может читать Deployment и StatefulSet "
            "(роль view это позволяет, а читать Secret — нет), а также Argo CD, Helm-релизы и git.")
    if open_places:
        risk += " Секреты в командах запуска и строках подключения видны открытым текстом даже в отчёте kubescape."
    if masked_env:
        risk += (" Значения переменных окружения kubescape скрыл, но видно, что они заданы строкой, а не "
                 "ссылкой на Secret" + (f": {', '.join(examples)}." if examples else "."))
    elif examples:
        risk += f" Среди них: {', '.join(examples)}."
    action = (f"Открытые секреты перенести в {store} и сменить: они уже разошлись по копиям. " if open_places else "")
    if masked_env:
        action += "Для переменных со скрытым значением проверить в чартах, что там записано; настоящие секреты — туда же. "
    action += "Список мест — лист «Секреты»."
    secrets_rng = "'Секреты'!$A$2:$A$5000"
    plan = [
        ("P1", "Секреты в описаниях подов (нашёл скрипт разбора, не kubescape)", risk, action,
         "не из отчёта kubescape: места нашёл ks_sanitize.py по именам и шаблонам; у kubescape "
         "похожая проверка C-0012, в набор CIS она не входит",
         f"=(COUNTA({secrets_rng})-COUNTIF({secrets_rng},\"{ENV_MASKED}\"))&\" открытым текстом, \"&"
         f"COUNTIF({secrets_rng},\"{ENV_MASKED}\")&\" переменных со скрытым значением\"",
         "команды сервисов, ИБ"),
        ("P1", "Pod Security Admission не включён в прикладных namespace",
         "В таких namespace можно запустить привилегированный под, примонтировать файловую систему "
         "узла или занять его сеть и порты — это прямой путь к захвату узла.",
         "Метки pod-security.kubernetes.io на namespace: сначала warn и audit на уровне baseline, "
         "найти нарушителей (kubectl label --dry-run=server), затем enforce=baseline; для прикладных "
         "namespace — restricted.",
         "C-0193, C-0197–C-0201, C-0203, C-0204",
         f"={failed('C-0193')}&\" из \"&{total('C-0193')}&\" namespace\"", "платформа"),
        ("P1", "Роль cluster-admin у личных и сервисных аккаунтов, широкие права RBAC",
         "Полный доступ к кластеру, включая все Secret. Токены сервисных аккаунтов не привязаны к "
         "SSO и второму фактору и легко утекают.",
         "Людям — вход через OIDC с ограниченными ролями вместо личных сервисных аккаунтов; "
         "cluster-admin оставить аварийной учётке; CI и Argo CD — права только на свои namespace. "
         "Большинство RBAC-находок закроется вместе с этим. Список — лист «RBAC».",
         "C-0185, C-0186, C-0188, C-0191, C-0278–C-0282",
         f"=COUNTIFS('RBAC'!$D$2:$D$2000,\"человек\",'RBAC'!$G$2:$G$2000,\"да\")&\" привязок cluster-admin "
         f"к людям; всего субъектов с cluster-admin: \"&{failed('C-0185')}", "платформа"),
        ("P2", "Нет сетевых политик",
         "Любой под может обратиться к любому: взлом одного сервиса открывает путь к базам и "
         "сервисам с ПДн.",
         "Запрет входящего трафика по умолчанию и явные разрешения, начиная с namespace с ПДн."
         + (" Сначала проверить, нет ли политик CiliumNetworkPolicy: kubescape их не видит." if cilium else ""),
         "C-0206", f"={failed('C-0206')}&\" из \"&{total('C-0206')}&\" namespace\"",
         "платформа, команды сервисов"),
        ("P2", "Поды без securityContext и seccomp, с токеном сервисного аккаунта",
         "При взломе приложения атакующий работает от root, может повышать привилегии и сразу "
         "получает токен для API Kubernetes.",
         "Значения по умолчанию в общих шаблонах Helm-чартов: runAsNonRoot, allowPrivilegeEscalation: "
         "false, readOnlyRootFilesystem, capabilities.drop: [ALL], seccompProfile: RuntimeDefault, "
         "automountServiceAccountToken: false. Одна правка шаблона закрывает сотни находок. "
         "Внедрять по namespace: часть приложений может не запуститься.",
         "C-0211, C-0210, C-0190, C-0189",
         f"={failed('C-0211')}&\" из \"&{total('C-0211')}&\" нагрузок\"", "команды сервисов"),
        ("P2", "Настройки API-сервера",
         "API-сервер не проверяет сертификаты kubelet (возможен перехват exec и logs), не запрещает "
         "externalIPs у Service (CVE-2020-8554), не ограничивает шифры TLS.",
         ("Менять через переменные Kubespray, иначе ручные правки манифестов затрутся. " if kubespray else "")
         + "--kubelet-certificate-authority вместе с серверными сертификатами kubelet; admission-плагины "
         "DenyServiceExternalIPs и EventRateLimit; --tls-cipher-suites без RC4 и CBC (значение из отчёта "
         "не копировать); --service-account-extend-token-expiration=false.",
         "C-0117, C-0121, C-0123, C-0134, C-0277, C-0283, C-0290",
         f"={failed('C-0117')}&\" узла control plane\"", "платформа"),
        ("P3", "Ротация журнала аудита API-сервера",
         "Аудит включён, но архивов меньше и файлы короче рекомендованного: истории событий может не "
         "хватить для расследования.",
         "--audit-log-maxbackup не меньше 10 и --audit-log-maxsize не меньше 100 либо отправка аудита в "
         "централизованное хранилище с нужным сроком хранения.",
         "C-0132, C-0133", f"={failed('C-0132')}&\" узла control plane\"", "платформа"),
        ("P3", "Secret через переменные окружения",
         "Значения видны в /proc, дампах и иногда в логах.",
         "Монтировать Secret файлом.", "C-0207", f"={failed('C-0207')}&\" нагрузок\"", "команды сервисов"),
        ("инфо", "Системные поды с сетью и процессами узла",
         "Все такие поды системные: им это нужно." if host_all_platform else
         "Среди подов с сетью или процессами узла есть прикладные — см. лист «Находки».",
         "Оформить исключения в kubescape, чтобы не мешали в следующих сканах.",
         "C-0041, C-0275", f"={failed('C-0041')}&\" с hostNetwork, \"&{failed('C-0275')}&\" с hostPID\"",
         "платформа"),
        ("инфо", "Часть проверок CIS не выполнялась",
         "Файлы и настройки на узлах, etcd, kubelet: CLI kubescape их не проверяет, итоговая оценка "
         "их не учитывает.",
         "Разово запустить kube-bench на узлах control plane и рабочих узлах или поставить оператор kubescape.",
         "лист «Проверки», статус «не проверялась»",
         f"=COUNTIF({CHECKS}!$E$2:$E$500,\"не проверялась\")&\" из \"&COUNTA({CHECKS}!$A$2:$A$500)&\" проверок\"",
         "платформа, ИБ"),
        ("инфо", "Ложные и ручные пункты",
         "C-0209 — ручной пункт, kubescape выносит все namespace на просмотр. C-0202 — Windows-контейнеры. "
         "C-0212 — см. лист «Проверки».",
         "Работ не требует.", "C-0209, C-0202, C-0212", None, "—"),
    ]
    sheet = Sheet(wb, "План", ["№", "Приоритет", "Проблема", "Чем опасно", "Что сделать",
                               "Закрывает проверки", "Масштаб", "Кто"],
                  [4, 12, 34, 55, 70, 24, 26, 18], wrap=(3, 4, 5, 6, 7), priority_col=2, restore=restore)
    for i, row in enumerate(plan, 1):
        sheet.add([i, *row])
    sheet.finish()

    # --- Проверки: все пункты CIS из отчёта
    sheet = Sheet(wb, "Проверки", ["ID", "Пункт CIS", "Название", "Критичность", "Статус", "Приоритет",
                                   "Не прошли", "Прошли (с исключениями)", "Из них исключения",
                                   "Что не так", "Что сделать", "Кто"],
                  [9, 8, 55, 12, 14, 10, 10, 12, 11, 60, 70, 18], wrap=(3, 10, 11), priority_col=6,
                  restore=restore)
    rows = []
    for c in data["controls"]:
        if c["status"] == "failed":
            note = notes.get(c["id"], ("P3", "платформа", "", ""))
        elif c["status"] == "skipped":
            note = skipped_notes.get(c["id"], SKIPPED_NOTE)
        else:
            note = ("—", "", "", "")
        num, title = cis_parts(c["name"])
        sev = SEVERITY_RU.get(c["severity"], c["severity"])
        status = STATUS_RU.get(c["status"], c["status"])
        rows.append([c["id"], num, title, sev, status, note[0], c["failed"], c["passed"], c["ignored"],
                     note[2], note[3], note[1]])
    rows.sort(key=lambda r: (STATUS_ORDER.get(r[4], 3), PRIORITY_ORDER.get(r[5], 5),
                             SEVERITY_ORDER.get(r[3], 4), -r[6], r[0]))
    for row in rows:
        sheet.add(row)
    sheet.finish()
    control_info = {r[0]: r for r in rows}

    # --- Находки: каждая пара «ресурс — непройденная проверка»
    sheet = Sheet(wb, "Находки", ["Зона", "Namespace", "Тип", "Имя", "Проверка", "Пункт CIS",
                                  "Название проверки", "Критичность", "Приоритет", "Что исправить",
                                  "Роль", "Привязка"],
                  [11, 20, 14, 40, 9, 8, 46, 12, 10, 70, 30, 34], wrap=(7, 10), priority_col=9,
                  restore=restore)
    rows = []
    for f in failed_findings:
        ns = f["name"] if f["kind"] == "Namespace" else f["namespace"]
        info = control_info.get(f["control"], [f["control"], "", "", "", "", "P3"])
        rows.append([zone_of(ns), ns, f["kind"], f["name"], f["control"], info[1], info[2], info[3],
                     info[5], fix_text(f), f.get("role", ""), f.get("binding", "")])
    rows.sort(key=lambda r: (r[0], r[1], r[2], r[3], PRIORITY_ORDER.get(r[8], 5), r[4]))
    for row in rows:
        sheet.add(row)
    sheet.finish()

    # --- RBAC: субъекты и роли с опасными правами
    sheet = Sheet(wb, "RBAC", ["Тип", "Namespace", "Имя", "Чей", "Роль", "Привязка", "cluster-admin",
                               "Не пройдено проверок", "Проверки"],
                  [15, 18, 34, 22, 34, 44, 13, 12, 60], wrap=(6, 9), restore=restore)
    grouped = {}
    for f in failed_findings:
        if f["kind"] not in RBAC_KINDS:
            continue
        key = (f["kind"], f["namespace"], f["name"], f.get("role", ""), f.get("binding", ""))
        grouped.setdefault(key, [f, set()])[1].add(f["control"])
    rows = []
    for (kind, ns, name, role, binding), (f, controls) in grouped.items():
        admin = "да" if role.endswith(" cluster-admin") else ""
        rows.append([kind, ns, name, whose(f), role, binding, admin, len(controls), ", ".join(sorted(controls))])
    rows.sort(key=lambda r: (r[6] != "да", -r[7], r[0], r[2]))
    for row in rows:
        sheet.add(row)
    sheet.finish()

    # --- Секреты: где в оригинале лежат пароли и токены, без значений
    sheet = Sheet(wb, "Секреты", ["Вид", "Зона", "Namespace", "Тип", "Ресурс", "Где", "Что сделать"],
                  [22, 11, 20, 12, 34, 50, 70], wrap=(6, 7), restore=restore)
    for kind_of_place, ns, kind, name, detail in sorted(secret_places, key=lambda p: (p[1], p[3], p[0], p[4])):
        sheet.add([kind_of_place, zone_of(ns), ns, kind, name, detail, SECRET_ACTION.get(kind_of_place, "")])
    sheet.finish()

    # --- Сводка
    fw = (data.get("frameworks") or [{}])[0].get("name", "")
    fw_title = re.sub(r"^cis-v", "CIS Kubernetes Benchmark v", fw)
    ws = summary
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 80
    ws["A1"] = f"Разбор отчёта kubescape — кластер {src.parent.name}"
    ws["A1"].font = TITLE
    ws["A2"] = (f"{fw_title} · скан {data.get('generationTime', '')[:10]} · "
                f"Kubernetes {data.get('kubernetes', '')} · файл {base}.json")
    ws["A3"] = ("Имена людей настоящие: файл только для внутреннего пользования." if names else
                "Таблица собрана из выжимки без секретов (ks_sanitize.py). Имена людей заменены "
                "псевдонимами user-NN, соответствие — в файле users.csv у владельца отчёта.")
    for cell in ("A2", "A3"):
        ws[cell].font = FONT
    r = 5
    for col, text in enumerate(["Показатель", "Значение", "Комментарий"], 1):
        cell = ws.cell(row=r, column=col, value=text)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
    findings_rng = "'Находки'!$A$2:$A$10000"
    indicators = [
        ("Оценка kubescape по CIS, %", data.get("complianceScore"),
         "из отчёта; считается только по выполненным проверкам"),
        ("Проверок в отчёте", f"=COUNTA({CHECKS}!$A$2:$A$500)", ""),
        ("   не пройдено", f"=COUNTIF({CHECKS}!$E$2:$E$500,\"не пройдена\")", "лист «Проверки»"),
        ("   пройдено", f"=COUNTIF({CHECKS}!$E$2:$E$500,\"пройдена\")", ""),
        ("   не проверялось", f"=COUNTIF({CHECKS}!$E$2:$E$500,\"не проверялась\")",
         "нужен доступ к узлам: kube-bench или оператор kubescape"),
        ("Находок (ресурс × проверка)", f"=COUNTA({findings_rng})", "лист «Находки»"),
        ("   в namespace платформы", f"=COUNTIF({findings_rng},\"платформа\")", ""),
        ("   в namespace сервисов", f"=COUNTIF({findings_rng},\"сервисы\")", ""),
        ("   на уровне кластера (RBAC)", f"=COUNTIF({findings_rng},\"кластер\")", ""),
        ("Похоже на секреты открытым текстом", f"=COUNTA({secrets_rng})-COUNTIF({secrets_rng},\"{ENV_MASKED}\")",
         "нашёл скрипт разбора, не kubescape: флаги и строки подключения в командах запуска; "
         "видны и в отчёте, в таблицу значения не попали"),
        ("Переменных с именами секретов, заданных строкой", f"=COUNTIF({secrets_rng},\"{ENV_MASKED}\")",
         "нашёл скрипт разбора, не kubescape; значения kubescape скрыл (XXXXXX), что там на самом деле — "
         "видно только в чартах"),
        ("Namespace без Pod Security Admission", f"={failed('C-0193')}", f"=\"из \"&{total('C-0193')}"),
        ("Namespace без сетевых политик", f"={failed('C-0206')}", f"=\"из \"&{total('C-0206')}"),
        ("Нагрузок без securityContext", f"={failed('C-0211')}", f"=\"из \"&{total('C-0211')}"),
        ("Субъектов RBAC с cluster-admin", f"={failed('C-0185')}", "лист «RBAC»"),
    ]
    for label, value, comment in indicators:
        r += 1
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=value)
        ws.cell(row=r, column=3, value=comment)
    r += 2
    for col, text in enumerate(["Что делать (подробно — лист «План»)", "Приоритет", "Масштаб"], 1):
        cell = ws.cell(row=r, column=col, value=text)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
    for i in range(len(plan)):
        r += 1
        ws.cell(row=r, column=1, value=f"='План'!C{i + 2}")
        ws.cell(row=r, column=2, value=f"='План'!B{i + 2}")
        ws.cell(row=r, column=3, value=f"=IF('План'!G{i + 2}=\"\",\"\",'План'!G{i + 2})")
    r += 2
    ws.cell(row=r, column=1, value="Как читать").font = BOLD
    legend = [
        "P1 — делать в первую очередь, P2 — следом, P3 — по возможности, «инфо» — работ не требует.",
        "Зона: «платформа» — namespace кластерной инфраструктуры (чинит команда кластера), «сервисы» — "
        "прикладные namespace (чинят их команды), «кластер» — объекты без namespace. Список namespace "
        "платформы задан в ks_report.py по названиям, его можно поправить.",
        "«Прошли» в kubescape включает ресурсы из исключений: например, 4 прошли и 4 в исключениях — "
        "значит, по-настоящему не прошёл никто.",
        "Непроверенные пункты — не «пройдены»: kubescape без доступа к узлам их пропускает.",
        "Лист «Секреты» и первый пункт плана — не находки kubescape: их нашёл скрипт разбора по именам "
        "переменных и флагов (PASSWORD, TOKEN, SECRET) и шаблонам вроде password=… Значения переменных "
        "окружения kubescape скрыл (XXXXXX), поэтому про них известно только, что они заданы строкой.",
    ]
    for text in legend:
        r += 1
        ws.cell(row=r, column=1, value=text)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    for row in ws.iter_rows(min_row=4):
        for cell in row:
            if cell.font != HEADER_FONT and cell.coordinate not in ("A1",):
                if not cell.font.bold:
                    cell.font = FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.sheet_view.showGridLines = False
    fit_to_width(ws)

    wb.save(dst)
    print(f"Таблица: {dst}")
    print(f"  проверок {len(data['controls'])}, находок {len(failed_findings)}, "
          f"строк RBAC {len(grouped)}, мест с секретами {len(secret_places)}")


if __name__ == "__main__":
    main()
